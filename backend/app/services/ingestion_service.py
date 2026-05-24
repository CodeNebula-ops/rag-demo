import os
import uuid
from pathlib import Path

import chardet
import fitz
import structlog
from docx import Document as DocxDocument
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.core.exceptions import DocumentProcessingError, FileValidationError
from app.core.text_preprocessor import normalize_text, remove_headers_footers
from app.models.document import Document, DocumentChunk
from app.services.chunking_service import chunk_document
from app.services.embedding_service import embed_texts
from app.services.vector_store_service import upsert_vectors
from app.services.audit_service import log_event
from app.utils.file_utils import ALLOWED_EXTENSIONS

logger = structlog.get_logger()


async def ingest_document(
    file_path: str,
    filename: str,
    file_size: int,
    session: AsyncSession,
) -> Document:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise FileValidationError(f"Unsupported file type: {ext}")

    if file_size > settings.max_file_size_mb * 1024 * 1024:
        raise FileValidationError(f"File exceeds {settings.max_file_size_mb}MB limit")

    title = Path(filename).stem.replace("_", " ").replace("-", " ").title()
    doc = Document(
        id=uuid.uuid4(),
        title=title,
        filename=filename,
        file_type=ext.lstrip("."),
        file_size_bytes=file_size,
        storage_path=file_path,
        status="processing",
    )
    session.add(doc)
    await session.flush()

    try:
        raw_text = _extract_text(file_path, ext)
        cleaned_text = normalize_text(raw_text)

        if not cleaned_text or len(cleaned_text.strip()) < 10:
            raise DocumentProcessingError("No text content extracted from document")

        chunks = chunk_document(cleaned_text, title)

        if not chunks:
            raise DocumentProcessingError("No chunks generated from document")

        chunk_texts = [c["chunk_text"] for c in chunks]
        embeddings = embed_texts(chunk_texts)

        payloads = []
        db_chunks = []

        for i, chunk_data in enumerate(chunks):
            chunk_id = uuid.uuid4()
            payloads.append({
                "document_id": str(doc.id),
                "document_title": title,
                "chunk_text": chunk_data["chunk_text"],
                "section_path": chunk_data.get("section_path", ""),
                "chunk_index": chunk_data["chunk_index"],
                "language": chunk_data.get("language", "en"),
                "is_active": True,
            })

            db_chunks.append(DocumentChunk(
                id=chunk_id,
                document_id=doc.id,
                chunk_index=chunk_data["chunk_index"],
                chunk_text=chunk_data["chunk_text"],
                section_path=chunk_data.get("section_path"),
                chunk_level=chunk_data.get("chunk_level", "paragraph"),
                token_count=chunk_data.get("token_count", 0),
                language=chunk_data.get("language", "en"),
            ))

        point_ids = upsert_vectors(embeddings.tolist(), payloads)

        for i, db_chunk in enumerate(db_chunks):
            db_chunk.qdrant_point_id = point_ids[i]

        session.add_all(db_chunks)

        doc.total_chunks = len(db_chunks)
        doc.status = "active"
        await session.flush()

        from app.services.retrieval_service import add_to_bm25_index
        add_to_bm25_index(chunk_texts, point_ids)

        await log_event(
            session,
            event_type="DOCUMENT_UPLOADED",
            resource_type="document",
            resource_id=str(doc.id),
            detail={"filename": filename, "chunks": len(db_chunks)},
        )

        logger.info("document_ingested", document_id=str(doc.id), chunks=len(db_chunks))
        return doc

    except (FileValidationError, DocumentProcessingError):
        raise
    except Exception as e:
        doc.status = "failed"
        await session.flush()
        logger.error("ingestion_failed", document_id=str(doc.id), error=str(e))
        raise DocumentProcessingError(str(e))


def _extract_text(file_path: str, ext: str) -> str:
    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return _extract_docx(file_path)
    elif ext in (".txt", ".md"):
        return _extract_text_file(file_path)
    raise FileValidationError(f"Cannot extract text from {ext}")


def _extract_pdf(file_path: str) -> str:
    doc = fitz.open(file_path)
    pages = []
    for page in doc:
        text = page.get_text()
        pages.append(text)
    doc.close()

    pages = remove_headers_footers(pages)
    return "\n\n".join(pages)


def _extract_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    parts = []
    for para in doc.paragraphs:
        if para.style and para.style.name.startswith("Heading"):
            level = para.style.name.replace("Heading ", "")
            prefix = "#" * min(int(level) if level.isdigit() else 1, 3)
            parts.append(f"{prefix} {para.text}")
        elif para.text.strip():
            parts.append(para.text)
    return "\n\n".join(parts)


def _extract_text_file(file_path: str) -> str:
    with open(file_path, "rb") as f:
        raw = f.read()
    detected = chardet.detect(raw)
    encoding = detected.get("encoding", "utf-8") or "utf-8"
    return raw.decode(encoding, errors="replace")
